"""Makes the before-and-after check's scoring spreadsheet (docs/pilot/pilot-check-scoring.xlsx).

The team types each learner's answers into it, by code, after the first and the
last session; it works out the scores and the change. The answer key comes from
content/pilot-check/en.json, so run this again whenever the questions change:

    python3 scripts/pilot_check_sheet.py [output.xlsx]

docs/notes/pilot-check.md says how the check is run and scored.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
CHECK = json.loads((ROOT / "content/pilot-check/en.json").read_text())
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs/pilot/pilot-check-scoring.xlsx"

QUESTIONS = CHECK["questions"]
N = len(QUESTIONS)
TAUGHT = sum(1 for q in QUESTIONS if q["taught"])
UNTAUGHT = N - TAUGHT
LETTERS = "ABC"
ANSWER_ROWS = 200  # learners × 2 sessions
RESULT_ROWS = 100  # learners

FONT = "Arial"
BODY = Font(name=FONT, size=11)
BOLD = Font(name=FONT, size=11, bold=True)
TITLE = Font(name=FONT, size=14, bold=True)
GREY = Font(name=FONT, size=11, color="595959")
INPUT = PatternFill("solid", fgColor="FFFF99")  # yellow: type here
FIXED = PatternFill("solid", fgColor="E7E6E6")  # grey: from the key, don't change
HEAD = PatternFill("solid", fgColor="D9E1F2")
THIN = Side(style="thin", color="A6A6A6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")


def style_all(ws):
    for row in ws.iter_rows():
        for cell in row:
            if cell.font == Font():
                cell.font = BODY


wb = Workbook()

# ── How to use ───────────────────────────────────────────────────────────────
how = wb.active
how.title = "How to use"
how.column_dimensions["A"].width = 110
lines = [
    ("Before-and-after check: scoring", TITLE),
    ("For the Thinkerwell team only. Learners are identified by their pilot code, never a name.", BOLD),
    ("", BODY),
    ("1. After each session, type one row per learner on the Answers sheet: the code, First or Last, the date,", BODY),
    (f"   the letter each learner circled for questions 1 to {N} (A, B or C; leave it empty if they circled none),", BODY),
    ("   both team members' writing marks (0 to 6 each, from the rubric on the answer key) and the five faces", BODY),
    ("   (1 = the first face, \"Not yet\", to 4 = the last, \"Yes\"). Only the yellow cells are for typing.", BODY),
    ("2. On the Results sheet, type each learner's code once in the yellow column. Everything else fills in.", BODY),
    ("3. Grey cells come from the answer key (the Key sheet). Don't change them.", BODY),
    ("", BODY),
    ("What the numbers mean", BOLD),
    (f"Score: correct answers out of {N}. Taught: out of the {TAUGHT} questions from lessons the pilot teaches.", BODY),
    (f"Untaught: out of the {UNTAUGHT} questions from lessons it doesn't. If scores rise on taught questions but not on", BODY),
    ("untaught ones, the lessons are the likely reason, not just taking the check twice.", BODY),
    ("Normalised gain: (after − before) ÷ (16 − before). It is fair to learners who start high. Empty when before is 16.", BODY),
    ("Writing: the average of the two markers' marks, out of 6. \"Check\" means the markers differ by 2 or more:", BODY),
    ("talk it through and agree, then type the agreed mark in both columns.", BODY),
    ("Confidence: the average of the five faces, from 1 to 4.", BODY),
    ("The summary at the top of Results counts only learners with both a first and a last session.", BODY),
    ("", BODY),
    ("Keep this file private. Delete it, with the paper sheets, within 6 months of the pilot ending (the consent form).", BOLD),
    ("", BODY),
    ("An example row on the Answers sheet looks like this (don't type it in):", BOLD),
]
for i, (text, font) in enumerate(lines, start=1):
    cell = how.cell(row=i, column=1, value=text)
    cell.font = font
example_row = len(lines) + 1
example = ["HLP-07", "First", "2026-10-20"] + ["A", "C", "B", "B", "A", "C", "B", "A", "B", "A", "C", "C", "A", "B", "C", "B"] + [4, 3, 2, 3, 3, 2, 4]
headers = ["Code", "Session", "Date"] + [f"Q{i}" for i in range(1, N + 1)] + ["Writing 1", "Writing 2", "Face 1", "Face 2", "Face 3", "Face 4", "Face 5"]
for col, (head, value) in enumerate(zip(headers, example), start=1):
    how.cell(row=example_row, column=col + 1, value=head).font = BOLD
    how.cell(row=example_row + 1, column=col + 1, value=value).font = GREY
how.cell(row=example_row, column=1, value="")

# ── Key ──────────────────────────────────────────────────────────────────────
key = wb.create_sheet("Key")
for col, head in enumerate(["Question", "Correct answer", "Lesson", "Taught in the pilot", "Area", "Question (English)"], start=1):
    cell = key.cell(row=1, column=col, value=head)
    cell.font, cell.fill, cell.border = BOLD, HEAD, BOX
for i, q in enumerate(QUESTIONS, start=2):
    letter = LETTERS[[o["correct"] for o in q["options"]].index(True)]
    values = [i - 1, letter, q["lesson"], "Yes" if q["taught"] else "No", q["area"], q["question"]]
    for col, value in enumerate(values, start=1):
        cell = key.cell(row=i, column=col, value=value)
        cell.font, cell.border = BODY, BOX
key.cell(row=N + 3, column=1, value="Source: content/pilot-check/en.json (made by scripts/pilot_check_sheet.py).").font = GREY
for col, width in zip("ABCDEF", (10, 15, 9, 18, 14, 70)):
    key.column_dimensions[col].width = width

# ── Answers ──────────────────────────────────────────────────────────────────
ans = wb.create_sheet("Answers")
Q0 = 4  # column D is Q1
QCOLS = [get_column_letter(Q0 + i) for i in range(N)]
W1, W2 = get_column_letter(Q0 + N), get_column_letter(Q0 + N + 1)
F0 = Q0 + N + 2
FCOLS = [get_column_letter(F0 + i) for i in range(5)]
SCORE, TAUGHT_C, UNTAUGHT_C, WRITE, CONF, MARKERS = (get_column_letter(F0 + 5 + i) for i in range(6))
first_q, last_q = QCOLS[0], QCOLS[-1]

heads = ["Code", "Session", "Date"] + [f"Q{i}" for i in range(1, N + 1)] + ["Writing 1 (0–6)", "Writing 2 (0–6)"] + [f"Face {i} (1–4)" for i in range(1, 6)] + [
    f"Score (of {N})", f"Taught (of {TAUGHT})", f"Untaught (of {UNTAUGHT})", "Writing (of 6)", "Confidence (1–4)", "Markers"]
for col, head in enumerate(heads, start=1):
    cell = ans.cell(row=1, column=col, value=head)
    cell.font, cell.fill, cell.border, cell.alignment = BOLD, HEAD, BOX, Alignment(wrap_text=True, horizontal="center", vertical="center")
ans.row_dimensions[1].height = 32

# Row 2: the right answers, row 3: taught (1) or not (0), both from the Key sheet.
ans.cell(row=2, column=1, value="Right answer").font = GREY
ans.cell(row=3, column=1, value="Taught (1) or not (0)").font = GREY
for i, col in enumerate(QCOLS):
    r = i + 2
    for row, formula in ((2, f"=Key!$B${r}"), (3, f'=IF(Key!$D${r}="Yes",1,0)')):
        cell = ans[f"{col}{row}"]
        cell.value, cell.fill, cell.font, cell.alignment, cell.border = formula, FIXED, GREY, CENTER, BOX

FIRST, LAST = 4, 3 + ANSWER_ROWS
for r in range(FIRST, LAST + 1):
    for col in range(1, F0 + 5):
        cell = ans.cell(row=r, column=col)
        cell.fill, cell.border = INPUT, BOX
        if col > 3:
            cell.alignment = CENTER
    qs = f"{first_q}{r}:{last_q}{r}"
    right = f"{first_q}$2:{last_q}$2"
    taught = f"{first_q}$3:{last_q}$3"
    blank = f'COUNTA(A{r},{qs})=0'
    formulas = {
        SCORE: f'=IF({blank},"",SUMPRODUCT(--({qs}={right})))',
        TAUGHT_C: f'=IF({blank},"",SUMPRODUCT(({qs}={right})*{taught}))',
        UNTAUGHT_C: f'=IF({blank},"",{SCORE}{r}-{TAUGHT_C}{r})',
        WRITE: f'=IF(COUNT({W1}{r}:{W2}{r})=0,"",AVERAGE({W1}{r}:{W2}{r}))',
        CONF: f'=IF(COUNT({FCOLS[0]}{r}:{FCOLS[-1]}{r})=0,"",AVERAGE({FCOLS[0]}{r}:{FCOLS[-1]}{r}))',
        MARKERS: f'=IF(COUNT({W1}{r}:{W2}{r})<2,"",IF(ABS({W1}{r}-{W2}{r})>=2,"Check",""))',
    }
    for col, formula in formulas.items():
        cell = ans[f"{col}{r}"]
        cell.value, cell.border, cell.alignment = formula, BOX, CENTER
        if col in (WRITE, CONF):
            cell.number_format = "0.0"
ans.freeze_panes = "B4"
ans.column_dimensions["A"].width = 12
ans.column_dimensions["B"].width = 10
ans.column_dimensions["C"].width = 12
for col in QCOLS:
    ans.column_dimensions[col].width = 5
for col in [W1, W2] + FCOLS + [SCORE, TAUGHT_C, UNTAUGHT_C, WRITE, CONF, MARKERS]:
    ans.column_dimensions[col].width = 11

rows = f"{FIRST}:{LAST}"
dv_session = DataValidation(type="list", formula1='"First,Last"', allow_blank=True, error="Type First or Last.")
dv_letter = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True, error="Type A, B or C, or leave it empty.")
dv_mark = DataValidation(type="whole", operator="between", formula1="0", formula2="6", allow_blank=True, error="A mark from 0 to 6.")
dv_face = DataValidation(type="whole", operator="between", formula1="1", formula2="4", allow_blank=True, error="1 to 4: the first face is 1.")
for dv in (dv_session, dv_letter, dv_mark, dv_face):
    ans.add_data_validation(dv)
dv_session.add(f"B{FIRST}:B{LAST}")
dv_letter.add(f"{first_q}{FIRST}:{last_q}{LAST}")
dv_mark.add(f"{W1}{FIRST}:{W2}{LAST}")
dv_face.add(f"{FCOLS[0]}{FIRST}:{FCOLS[-1]}{LAST}")
ans["A1"].comment = Comment("The learner's pilot code, as on their sheet (for example HLP-07). Never a name.", "Thinkerwell")
ans["B1"].comment = Comment("First (the first session) or Last (the last session).", "Thinkerwell")

# ── Results ──────────────────────────────────────────────────────────────────
res = wb.create_sheet("Results")
res["A1"] = "Before-and-after check: results"
res["A1"].font = TITLE
res["A2"] = "Type each learner's code once in column A (yellow), from row 15. The summary counts learners with both sessions."
res["A2"].font = GREY

R0, R1 = 15, 14 + RESULT_ROWS
cols = ["Code", f"Before (of {N})", f"After (of {N})", "Change", "Normalised gain", f"Taught before (of {TAUGHT})", f"Taught after (of {TAUGHT})",
        f"Untaught before (of {UNTAUGHT})", f"Untaught after (of {UNTAUGHT})", "Writing before (of 6)", "Writing after (of 6)",
        "Confidence before (1–4)", "Confidence after (1–4)", "Both sessions", "Note"]
for col, head in enumerate(cols, start=1):
    cell = res.cell(row=14, column=col, value=head)
    cell.font, cell.fill, cell.border, cell.alignment = BOLD, HEAD, BOX, Alignment(wrap_text=True, horizontal="center", vertical="center")
res.row_dimensions[14].height = 45


def pick(r, source_col, session):
    """The value in Answers' column for this code and session ("" when there is no row)."""
    found = f'COUNTIFS(Answers!$A${FIRST}:$A${LAST},$A{r},Answers!$B${FIRST}:$B${LAST},"{session}")'
    return (f'=IF(OR($A{r}="",{found}=0),"",SUMIFS(Answers!${source_col}${FIRST}:${source_col}${LAST},'
            f'Answers!$A${FIRST}:$A${LAST},$A{r},Answers!$B${FIRST}:$B${LAST},"{session}"))')


for r in range(R0, R1 + 1):
    res.cell(row=r, column=1).fill = INPUT
    pairs = {"B": (SCORE, "First"), "C": (SCORE, "Last"), "F": (TAUGHT_C, "First"), "G": (TAUGHT_C, "Last"),
             "H": (UNTAUGHT_C, "First"), "I": (UNTAUGHT_C, "Last"), "J": (WRITE, "First"), "K": (WRITE, "Last"),
             "L": (CONF, "First"), "M": (CONF, "Last")}
    for col, (source, session) in pairs.items():
        res[f"{col}{r}"] = pick(r, source, session)
    res[f"D{r}"] = f'=IF(N{r}=1,C{r}-B{r},"")'
    res[f"E{r}"] = f'=IF(OR(N{r}<>1,B{r}={N}),"",(C{r}-B{r})/({N}-B{r}))'
    res[f"N{r}"] = f'=IF(AND(B{r}<>"",C{r}<>""),1,0)'
    dupes = f'MAX(COUNTIFS(Answers!$A${FIRST}:$A${LAST},$A{r},Answers!$B${FIRST}:$B${LAST},"First"),COUNTIFS(Answers!$A${FIRST}:$A${LAST},$A{r},Answers!$B${FIRST}:$B${LAST},"Last"))'
    res[f"O{r}"] = (f'=IF($A{r}="","",IF({dupes}>1,"Two rows for one session: keep one",'
                    f'IF(B{r}="","No first session",IF(C{r}="","No last session",""))))')
    for col in range(1, 16):
        cell = res.cell(row=r, column=col)
        cell.border = BOX
        if col > 1:
            cell.alignment = CENTER
    for col in "JKLM":
        res[f"{col}{r}"].number_format = "0.0"
    res[f"E{r}"].number_format = "0.00"

rng = lambda col: f"{col}{R0}:{col}{R1}"  # noqa: E731
both = rng("N")
summary = [
    ("Learners with both sessions", f"=SUM({both})", "0"),
    (f"Average score before (of {N})", f'=IFERROR(AVERAGEIFS({rng("B")},{both},1),"")', "0.0"),
    (f"Average score after (of {N})", f'=IFERROR(AVERAGEIFS({rng("C")},{both},1),"")', "0.0"),
    ("Average change", f'=IFERROR(AVERAGEIFS({rng("D")},{both},1),"")', "+0.0;-0.0;0.0"),
    ("Share who improved", f'=IFERROR(COUNTIFS({both},1,{rng("D")},">0")/SUM({both}),"")', "0%"),
    ("Average normalised gain", f'=IFERROR(AVERAGEIFS({rng("E")},{both},1),"")', "0.00"),
    ("Taught questions: change in % right", f'=IFERROR((AVERAGEIFS({rng("G")},{both},1)-AVERAGEIFS({rng("F")},{both},1))/{TAUGHT},"")', "+0%;-0%;0%"),
    ("Untaught questions: change in % right", f'=IFERROR((AVERAGEIFS({rng("I")},{both},1)-AVERAGEIFS({rng("H")},{both},1))/{UNTAUGHT},"")', "+0%;-0%;0%"),
    ("Writing: average change (of 6)", f'=IFERROR(AVERAGEIFS({rng("K")},{both},1,{rng("J")},">=0",{rng("K")},">=0")-AVERAGEIFS({rng("J")},{both},1,{rng("J")},">=0",{rng("K")},">=0"),"")', "+0.0;-0.0;0.0"),
    ("Confidence: average change (1–4)", f'=IFERROR(AVERAGEIFS({rng("M")},{both},1,{rng("L")},">=1",{rng("M")},">=1")-AVERAGEIFS({rng("L")},{both},1,{rng("L")},">=1",{rng("M")},">=1"),"")', "+0.0;-0.0;0.0"),
]
for i, (label, formula, fmt) in enumerate(summary, start=3):
    res.cell(row=i, column=1, value=label).font = BOLD
    cell = res.cell(row=i, column=4, value=formula)
    cell.number_format, cell.font, cell.border = fmt, BOLD, BOX
res.column_dimensions["A"].width = 14
for col in "BCDEFGHIJKLMN":
    res.column_dimensions[col].width = 12
res.column_dimensions["O"].width = 32
res.freeze_panes = "B15"

for ws in (how, key, ans, res):
    style_all(ws)

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(f"Wrote {OUT.relative_to(ROOT) if OUT.is_relative_to(ROOT) else OUT}")
