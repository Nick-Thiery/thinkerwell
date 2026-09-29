"""Round trip for the review spreadsheet: export, edit like a reviewer, import.

Run: .venv/bin/python -m unittest discover -s scripts/i18n -p "test_*.py"
(npm run test:review does the same). Works on a copy of the repository's
content, so nothing real is changed.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", ".."))


class ReviewRoundTrip(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        for part in ["content", os.path.join("src", "i18n", "messages"), os.path.join("docs", "translation")]:
            shutil.copytree(os.path.join(REPO, part), os.path.join(self.tmp, part))
        self.env = {**os.environ, "THINKERWELL_ROOT": self.tmp}
        self.xlsx = os.path.join(self.tmp, "review.xlsx")
        self.run_script("export_review.py", "id", "--out", self.xlsx)

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def run_script(self, script, *args):
        return subprocess.run([sys.executable, os.path.join(HERE, script), *args], env=self.env,
                              capture_output=True, text=True)

    def load(self, *parts):
        with open(os.path.join(self.tmp, *parts), encoding="utf-8") as f:
            return json.load(f)

    def edit(self, changes):
        """changes: {row ID: {column name: value}}."""
        from openpyxl import load_workbook
        wb = load_workbook(self.xlsx)
        for ws in wb.worksheets:
            header = None
            for row in ws.iter_rows():
                values = [c.value for c in row]
                if header is None:
                    if values and values[0] == "ID":
                        header = values
                    continue
                for column, value in changes.get(values[0], {}).items():
                    row[header.index(column)].value = value
        wb.save(self.xlsx)

    def test_tabs_and_columns(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.xlsx)
        self.assertEqual(wb.sheetnames[:7], ["Interface", "Lesson 1", "Lesson 2", "Lesson 11", "Lesson 12", "Lesson 16", "Lesson 23"])
        self.assertEqual(wb.sheetnames[-4:], ["Check - History", "Check - Geography", "Check - Culture", "Check - Civics"])
        self.assertEqual(len(wb.sheetnames), 29)
        header = next(r for r in wb["Lesson 1"].iter_rows(values_only=True) if r[0] == "ID")
        self.assertEqual(list(header[:8]), ["ID", "Where it appears", "English", "Indonesian",
                                            "Back-translation into English", "Flag", "Reviewer", "Notes"])
        brief = " ".join(str(r[0]) for r in wb["Interface"].iter_rows(max_row=10, values_only=True) if r[0])
        for words in ("10 to 17", "kamu", "simpler", "Quick checks", "Fix rather than rewrite", "Notes"):
            self.assertIn(words, brief)

    def test_applies_changes_and_keeps_notes(self):
        self.edit({
            "L01:read.sections.0.heading": {"Indonesian": "Sumber adalah petunjuk masa lalu", "Reviewer": "RA"},
            "ui:nav.settings": {"Indonesian": "Setelan", "Notes": "Setelan lebih umum di HP?"},
            "L01:read.glossary.1.forms": {"Indonesian": "sumber-sumber, sumbernya"},
        })
        result = self.run_script("import_review.py", "id", self.xlsx)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        lesson = self.load("content", "id", "lessons", "L01.json")
        self.assertEqual(lesson["read"]["sections"][0]["heading"], "Sumber adalah petunjuk masa lalu")
        self.assertEqual(lesson["read"]["glossary"][1]["forms"], ["sumber-sumber", "sumbernya"])
        self.assertEqual(self.load("src", "i18n", "messages", "id.json")["nav"]["settings"], "Setelan")
        notes = self.load("docs", "translation", "id", "notes", "ui.json")
        self.assertEqual(notes["nav.settings"]["reviewerNote"], "Setelan lebih umum di HP?")
        self.assertEqual(self.load("docs", "translation", "id", "notes", "L01.json")["read.sections.0.heading"]["reviewer"], "RA")
        self.assertIn("3 change(s)", result.stdout)

    def test_applies_section_check_and_picture_changes(self):
        self.edit({
            "quiz-history:intro": {"Indonesian": "Dua belas pertanyaan tentang Pelajaran 1–9."},
            "L01:picture.text.8": {"Indonesian": "Catatan harian"},
            "ui:pages.course.lessonsBadge": {"Indonesian": "{count} pelajaran saja"},
        })
        result = self.run_script("import_review.py", "id", self.xlsx)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual(self.load("content", "id", "quizzes", "history.json")["intro"], "Dua belas pertanyaan tentang Pelajaran 1–9.")
        with open(os.path.join(self.tmp, "content", "id", "visuals", "L01.svg"), encoding="utf-8") as f:
            self.assertIn(">Catatan harian<", f.read())
        self.assertEqual(self.load("src", "i18n", "messages", "id.json")["pages"]["course"]["lessonsBadge"], {"other": "{count} pelajaran saja"})

    def test_refuses_broken_placeholders_and_empty_text(self):
        self.edit({
            "ui:lesson.number": {"Indonesian": "Pelajaran nomor"},
            "L01:title": {"Indonesian": ""},
        })
        result = self.run_script("import_review.py", "id", self.xlsx)
        self.assertIn("placeholders", result.stdout)
        self.assertIn("L01:title: the Indonesian is empty", result.stdout)
        with open(os.path.join(REPO, "src", "i18n", "messages", "id.json"), encoding="utf-8") as f:
            original = json.load(f)["lesson"]["number"]
        self.assertEqual(self.load("src", "i18n", "messages", "id.json")["lesson"]["number"], original)

    def test_skips_conflicts(self):
        path = os.path.join(self.tmp, "content", "id", "lessons", "L02.json")
        lesson = self.load("content", "id", "lessons", "L02.json")
        lesson["title"] = "Diubah di file"
        with open(path, "w", encoding="utf-8") as f:
            json.dump(lesson, f, ensure_ascii=False, indent=2)
        self.edit({"L02:title": {"Indonesian": "Diubah di spreadsheet"}})
        result = self.run_script("import_review.py", "id", self.xlsx)
        self.assertIn("1 conflict(s)", result.stdout)
        self.assertEqual(self.load("content", "id", "lessons", "L02.json")["title"], "Diubah di file")

    def test_dry_run_writes_nothing(self):
        self.edit({"L01:title": {"Indonesian": "Judul baru"}})
        before = self.load("content", "id", "lessons", "L01.json")
        result = self.run_script("import_review.py", "id", self.xlsx, "--dry-run")
        self.assertIn("1 change(s)", result.stdout)
        self.assertEqual(self.load("content", "id", "lessons", "L01.json"), before)


if __name__ == "__main__":
    unittest.main()
