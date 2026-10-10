#!/usr/bin/env python3
"""Apply reviewers' changes from a review spreadsheet (made by export_review.py).

Usage: import_review.py [<lang>] <file.xlsx> [--dry-run]   (<lang> is id when left out)

For every row, compares the reviewer's text in the translation column with
the hidden copy of what was exported:
- unchanged: nothing to do;
- changed: written into the translation file (src/i18n/messages/<lang>.json,
  content/<lang>/course.json, lessons/LNN.json or quizzes/<section>.json),
  or for a picture's words, into content/<lang>/visuals/LNN.svg (a label
  split over several lines there can't be changed from the sheet: the
  import says so, to be changed in the picture by hand);
- changed here AND in the file since the export: a conflict. It is listed
  and skipped, so nobody's fix is silently overwritten.

A change is refused (and listed) if it would break the file: an empty text,
a missing or extra {placeholder} in an interface message, or an unknown ID.
Reviewer names and notes are copied into docs/translation/<lang>/notes/*.json
(as "reviewer" and "reviewerNote") so questions aren't lost, and printed.

Afterwards, run `npm run check:content` and `npm test`: they check every key
word still appears in its reading, the simpler text is still simpler, and
nothing else changed.
"""
import argparse
import html
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
import translatable as T  # noqa: E402
import review_sheet as S  # noqa: E402

from openpyxl import load_workbook  # noqa: E402


def read_rows(path):
    """Every data row in every tab, as a dict keyed by column name."""
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        header = None
        for values in ws.iter_rows(values_only=True):
            if header is None:
                if values and values[0] == "ID":
                    header = [str(v).strip() if v is not None else "" for v in values]
                continue
            row = {name: values[i] if i < len(values) else None for i, name in enumerate(header) if name}
            if row.get("ID"):
                yield ws.title, row


def text(value):
    return "" if value is None else str(value).replace("\r\n", "\n").strip()


def to_value(path, cell):
    """The file value for a cell: glossary forms are a comma-separated list."""
    if path[-1] == "forms":
        return [part.strip() for part in cell.split(",") if part.strip()]
    return cell


def same(file_value, cell):
    return S.cell_text(file_value).strip() == cell


def main(argv):
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("lang_or_file", help="the language code (id, ms, vi), then the file; Indonesian when only the file is given")
    parser.add_argument("xlsx", nargs="?")
    parser.add_argument("--dry-run", action="store_true", help="report what would change, write nothing")
    args = parser.parse_args(argv)
    if args.xlsx is None:
        args.lang, args.xlsx = "id", args.lang_or_file
    else:
        args.lang = args.lang_or_file
    lang = args.lang
    column = S.language_name(lang)

    files, notes = {}, {}
    changes, conflicts, refused, questions, broken = [], [], [], [], []
    changed_files, changed_notes, changed_pictures = set(), set(), set()

    pictures = {}  # lesson name -> [path, svg text]

    def picture_for(name):
        if name not in pictures:
            files = S.picture_files(lang, name)
            if not files or not os.path.exists(files[1]):
                raise KeyError(name)
            with open(files[1], encoding="utf-8") as f:
                pictures[name] = [files[1], f.read()]
        return pictures[name]

    def data_for(name):
        if name not in files:
            eng_path, tr_path = S.translation_file(lang, name)
            if not os.path.exists(eng_path) or not os.path.exists(tr_path):
                raise KeyError(name)
            files[name] = (T.load(eng_path), T.load(tr_path), tr_path)
            notes[name] = S.load_notes(lang, name)
        return files[name]

    for tab, row in read_rows(args.xlsx):
        row_id = text(row.get("ID"))
        try:
            name, path = S.split_id(row_id)
            english, data, _ = data_for(name)
        except (KeyError, ValueError):
            refused.append(f"{tab}: unknown ID {row_id!r}")
            continue
        key = ".".join(str(p) for p in path)
        reviewer, note = text(row.get("Reviewer")), text(row.get("Notes"))
        if reviewer or note:
            entry = notes[name].setdefault(key, {})
            if reviewer:
                entry["reviewer"] = reviewer
            if note:
                entry["reviewerNote"] = note
                questions.append(f"{row_id} ({reviewer or 'no name'}): {note}")
            changed_notes.add(name)

        new, exported = text(row.get(column)), text(row.get(f"Exported {column} (do not edit)"))
        if new == exported:
            continue
        if path and path[0] == "picture":
            try:
                picture = picture_for(name)
            except KeyError:
                refused.append(f"{row_id}: lesson {name} has no translated picture")
                continue
            svg = picture[1]
            old_markup, new_markup = f">{html.escape(exported, quote=False)}<", f">{html.escape(new, quote=False)}<"
            if not new:
                refused.append(f"{row_id}: the {column} is empty")
            elif svg.count(old_markup) != 1:
                refused.append(f"{row_id}: \"{exported}\" isn't one piece of text in {os.path.relpath(picture[0], T.ROOT)} "
                               f"(split over lines, or changed since the export); change it in the picture by hand")
            else:
                picture[1] = svg.replace(old_markup, new_markup)
                changes.append(f"{row_id}: {exported!r} -> {new!r}")
                changed_pictures.add(name)
            continue
        current = S.ui_value(data, ".".join(str(p) for p in path)) if name == "ui" else T.get_path(data, path)
        if not same(current, exported) and not same(current, new):
            conflicts.append(f"{row_id}: changed in the file since the export; kept the file's text. "
                             f"Sheet: {new!r} | file now: {S.cell_text(current)!r}")
            continue
        if path[-1] != "forms" and not new:
            refused.append(f"{row_id}: the {column} is empty")
            continue
        if name == "ui":
            eng_text = S.ui_value(english, ".".join(path))
            if S.placeholders(new) != S.placeholders(eng_text):
                refused.append(f"{row_id}: placeholders {S.placeholders(new)} don't match the English {S.placeholders(eng_text)}")
                continue
            parent = T.get_path(data, path[:-1])
            if S.is_plural(parent.get(path[-1])):
                parent[path[-1]] = {"other": new}
            else:
                parent[path[-1]] = new
        else:
            value = to_value(path, new)
            if path[-1] == "forms" and not value:
                T.get_path(data, path[:-1]).pop("forms", None)
            else:
                T.set_path(data, path, value, english if path[-1] != "forms" else data)
        changes.append(f"{row_id}: {S.cell_text(current)!r} -> {new!r}")
        changed_files.add(name)
        # The translation now matches today's English (check_translation.py warns when they drift).
        notes[name].setdefault(key, {})["source"] = (
            S.ui_value(english, key) if name == "ui"
            else list((T.get_path(english, path[:-1]) or {}).get("forms", [])) if path[-1] == "forms"
            else T.get_path(english, path))

    # The files must still have the right shape before anything is written.
    for name, (english, data, _) in files.items():
        if name == "ui":
            continue
        for problem in T.problems(english, data) + [f"{p}: not translated" for p in T.missing(english, data)]:
            broken.append(f"{name}: {problem}")

    print(f"{len(changes)} change(s), {len(conflicts)} conflict(s), {len(refused)} refused, {len(questions)} note(s) from reviewers.")
    for title, items in (("Changes", changes), ("Conflicts (skipped)", conflicts), ("Refused", refused),
                         ("Notes from reviewers", questions), ("Problems (nothing written)", broken)):
        if items:
            print(f"\n{title}:")
            for item in items:
                print(f"  {item}")

    if broken:
        # Something is wrong with a whole file, not one row: write nothing.
        print("\nNothing was written: fix the problems above first.")
        return 1
    if args.dry_run:
        print("\nDry run: nothing was written.")
        return 0
    for name in changed_files:
        with open(files[name][2], "w", encoding="utf-8") as f:
            f.write(S.dumps(files[name][1]))
    for name in changed_pictures:
        with open(pictures[name][0], "w", encoding="utf-8") as f:
            f.write(pictures[name][1])
    for name in changed_files | changed_notes:
        os.makedirs(os.path.dirname(S.notes_path(lang, name)), exist_ok=True)
        with open(S.notes_path(lang, name), "w", encoding="utf-8") as f:
            f.write(S.dumps(notes[name]))
    print("\nWritten. Now run: npm run check:content && npm test")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
