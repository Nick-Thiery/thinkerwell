#!/usr/bin/env python3
"""Write the review spreadsheet for a translation.

Usage: export_review.py [<lang>] [--out FILE]   (<lang> is id when left out)
  (default FILE: docs/translation/<lang>/review/thinkerwell-<lang>-review.xlsx)

One tab for the interface (with the brief for reviewers at the top), then
one tab per lesson (its words, then its picture's), priority lessons first
(see review_sheet.PRIORITY_LESSONS), then one tab per section check.
Columns: ID, where it appears, English, the translation, a back-translation
into English, Flag, Reviewer, Notes, and a hidden copy of the exported
translation that import_review.py uses to see what a reviewer changed.

Needs openpyxl (sh scripts/setup-python.sh installs it into .venv).
"""
import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(__file__))
import translatable as T  # noqa: E402
import review_sheet as S  # noqa: E402

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

BRIEF = [
    "How to review this translation",
    "Readers are 10 to 17 years old, many learning to read, some reading slowly. Most are refugee or displaced young people in Jakarta.",
    "Plain, everyday Indonesian. Speak to the learner as \"kamu\" (never \"Anda\"). Pages for adults (Educators, teacher guides, Settings, For organisations, the consent form) use \"Anda\". No \"silakan\" or \"tolong\", no exclamation marks. It should not sound stiff or like a textbook.",
    "The SIMPLER text must stay clearly simpler than the standard text: short sentences, everyday words.",
    "Quick checks: the \"Where it appears\" column says which option is the CORRECT answer. Check it is still clearly right in Indonesian and the others clearly wrong.",
    "Fix rather than rewrite: change only what needs changing, straight in the Indonesian column. Put your name or initials in Reviewer on rows you checked.",
    "Put questions and doubts in Notes. Rows with a Flag are where the meaning may have shifted or the translator was unsure: start with those.",
    "Names of people and places stay as they are. Decide the repeated key terms first (docs/translation/id/KEY_TERMS.md), so they are the same everywhere.",
    "Each lesson tab ends with the words in its picture: they must fit their place in the picture, so keep them short. The last four tabs are the section checks.",
    "Don't edit the ID column or the hidden last column, and don't add or delete rows; Nick's script reads your changes back in by ID.",
]

def brief_for(lang):
    """The reviewers' brief for a language: the Indonesian one is written out, the others are made from it."""
    if lang == "id":
        return BRIEF[1:]
    name = S.language_name(lang)
    lines = [line.replace("Indonesian", name).replace("docs/translation/id/", f"docs/translation/{lang}/") for line in BRIEF[1:]]
    if lang in MALAY_STYLE:
        lines[1] = MALAY_STYLE[lang]
    return lines


MALAY_STYLE = {
    "ms": "Plain, everyday Malaysian Malay (Bahasa Melayu as Malaysian schools use it, DBP spelling), not Indonesian. Speak to the learner as \"kamu\"; "
          "pages for adults (Educators, teacher guides, Settings, For organisations, the consent form) say \"anda\", in lower case. "
          "No \"sila\" or \"tolong\", no exclamation marks. An Indonesian word in the Malay is a mistake (for example gratis, bisa, mobil, kantor). "
          "It should not sound stiff or like a textbook.",
}


HEADER_FILL = PatternFill("solid", fgColor="1F1B24")
FLAG_FILL = PatternFill("solid", fgColor="F6DCC8")  # a light burnt orange (the "not quite" colour; never red)
EDIT_FILL = PatternFill("solid", fgColor="FFFFD6")  # the column reviewers edit
BRIEF_FILL = PatternFill("solid", fgColor="EDE7FB")  # lavender: help
WIDTHS = [26, 30, 55, 55, 55, 40, 14, 40, 10]
WRAP = Alignment(wrap_text=True, vertical="top")


def flag_text(entry):
    if not entry.get("flag"):
        return ""
    return entry.get("note", "").strip() or "Check this row."


def add_sheet(wb, title, rows, lang, brief=None):
    ws = wb.create_sheet(title)
    r = 1
    if brief:
        ws.cell(row=1, column=1, value=brief[0]).font = Font(bold=True, size=14)
        for i, line in enumerate(brief[1:], start=2):
            ws.cell(row=i, column=1, value=f"• {line}")
            ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
            ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[i].height = 32
        for i in range(1, len(brief) + 1):
            for c in range(1, 7):
                ws.cell(row=i, column=c).fill = BRIEF_FILL
        r = len(brief) + 2
    header_row = r
    for c, name in enumerate(S.columns(lang), start=1):
        cell = ws.cell(row=r, column=c, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = WRAP
    for row_id, where, english, translation, entry in rows:
        r += 1
        values = [row_id, where, S.cell_text(english), S.cell_text(translation), entry.get("back", ""),
                  flag_text(entry), "", "", S.cell_text(translation)]
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.alignment = WRAP
        ws.cell(row=r, column=4).fill = EDIT_FILL
        if values[5]:
            ws.cell(row=r, column=6).fill = FLAG_FILL
    for c, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.column_dimensions[get_column_letter(len(S.columns(lang)))].hidden = True
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(S.columns(lang)) - 1)}{r}"
    return ws


def export(lang, out):
    wb = Workbook()
    wb.remove(wb.active)
    language = S.language_name(lang)
    brief = [BRIEF[0] + f" ({language})"] + brief_for(lang)
    ui_rows = S.rows_for(lang, "ui") + S.rows_for(lang, "course")
    add_sheet(wb, "Interface", ui_rows, lang, brief)
    counts = {"Interface": len(ui_rows)}
    for name in S.lesson_names():
        rows = S.rows_for(lang, name)
        number = int(name[1:])
        title = f"Lesson {number}"
        add_sheet(wb, title, rows, lang)
        counts[title] = len(rows)
    for section in S.quiz_names():
        rows = S.rows_for(lang, f"quiz-{section}")
        title = S.QUIZ_TITLES.get(section, f"Check - {section}")
        add_sheet(wb, title, rows, lang)
        counts[title] = len(rows)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    flagged = sum(1 for name in ["ui", "course"] + S.lesson_names() + [f"quiz-{q}" for q in S.quiz_names()]
                  for _, _, _, _, e in S.rows_for(lang, name) if e.get("flag"))
    print(f"Wrote {os.path.relpath(out, T.ROOT)}: {len(counts)} tabs, {sum(counts.values())} rows, {flagged} flagged.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("lang", nargs="?", default="id", help="id (the default), ms or vi")
    parser.add_argument("--out")
    args = parser.parse_args()
    default = os.path.join(T.ROOT, "docs", "translation", args.lang, "review", f"thinkerwell-{args.lang}-review.xlsx")
    export(args.lang, os.path.abspath(args.out) if args.out else default)
